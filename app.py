import io
import calendar
import random
import re
from datetime import datetime
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
import streamlit as st

def obter_domingos(mes, ano):
    c = calendar.Calendar(firstweekday=calendar.SUNDAY)
    mes_calendario = c.monthdatescalendar(ano, mes)
    domingos = []
    for semana in mes_calendario:
        domingo = semana[0]
        if domingo.month == mes and domingo.year == ano:
            domingos.append({
                "dia": domingo.day,
                "data_str": domingo.strftime("%d/%m/%Y")
            })
    return domingos

def exportar_excel_completo(dados_escala, contagem_escalados, lista_texto_final, horarios_labels):
    wb = Workbook()
    
    # --- ABA 1: ESCALA VISUAL ---
    ws_escala = wb.active
    ws_escala.title = "Escala"
    ws_escala.views.sheetView[0].showGridLines = True

    # Título Principal (Mesclando baseado no número dinâmico de colunas)
    total_colunas = len(horarios_labels) + 1
    ws_escala.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_colunas)
    ws_escala["A1"] = "ESCALA DE VOLUNTÁRIOS"
    ws_escala["A1"].font = Font(name="Calibri", size=18, bold=True, color="2E7D32")
    ws_escala["A1"].alignment = Alignment(horizontal="center", vertical="center")

    headers = ["DATA"] + horarios_labels
    ws_escala.append([])  
    ws_escala.append(headers)  

    header_fill = PatternFill(start_color="555555", end_color="555555", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    center_align = Alignment(horizontal="center", vertical="center")

    for col_num in range(1, total_colunas + 1):
        cell = ws_escala.cell(row=3, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    thin_border = Side(style="thin", color="E0E0E0")
    border_style = Border(left=thin_border, right=thin_border, top=thin_border, bottom=thin_border)
    vago_fill = PatternFill(start_color="FFEBEE", end_color="FFEBEE", fill_type="solid")
    vago_font = Font(name="Calibri", size=10, bold=True, color="D32F2F")
    aviso_font = Font(name="Calibri", size=10, bold=True, color="E65100")

    for i, row_data in enumerate(dados_escala):
        row_num = i + 4
        
        # Monta a linha dinamicamente de acordo com as colunas ativas
        linha_excel = [row_data["DATA"]]
        for h in horarios_labels:
            linha_excel.append(row_data.get(h, "VAGO"))
            
        ws_escala.append(linha_excel)

        bg_color = "FFFFFF" if i % 2 == 0 else "F9F9F9"
        row_fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")

        for col_num in range(1, total_colunas + 1):
            cell = ws_escala.cell(row=row_num, column=col_num)
            cell.fill = row_fill
            cell.alignment = center_align
            cell.border = border_style
            cell.font = Font(name="Calibri", size=10)

            if col_num == 1:
                cell.font = Font(name="Calibri", size=10, bold=True)

            val_str = str(cell.value)
            if val_str == "VAGO":
                cell.fill = vago_fill
                cell.font = vago_font
            elif "⚠️" in val_str:
                cell.font = aviso_font

    # Ajusta larguras das colunas dinamicamente
    ws_escala.column_dimensions["A"].width = 18
    for col_idx in range(2, total_colunas + 1):
        letra_coluna = chr(64 + col_idx) if col_idx <= 26 else "B" # Fallback simples
        ws_escala.column_dimensions[letra_coluna].width = 35
        
    ws_escala.row_dimensions[1].height = 40
    for r in range(3, ws_escala.max_row + 1):
        ws_escala.row_dimensions[r].height = 25

    # --- ABA 2: RESUMO E ESTATÍSTICAS ---
    ws_resumo = wb.create_sheet(title="Resumo e Estatísticas")
    ws_resumo.views.sheetView[0].showGridLines = True
    
    ws_resumo["A1"] = "CONTROLE DE PARTICIPAÇÕES"
    ws_resumo["A1"].font = Font(name="Calibri", size=14, bold=True, color="1B5E20")
    
    ws_resumo["A3"] = "Voluntário"
    ws_resumo["B3"] = "Vezes Escalado"
    ws_resumo["A3"].font = Font(name="Calibri", bold=True, color="FFFFFF")
    ws_resumo["B3"].font = Font(name="Calibri", bold=True, color="FFFFFF")
    ws_resumo["A3"].fill = header_fill
    ws_resumo["B3"].fill = header_fill
    ws_resumo["A3"].alignment = center_align
    ws_resumo["B3"].alignment = center_align
    
    voluntarios_ordenados = sorted(contagem_escalados.items(), key=lambda item: item[1], reverse=True)
    
    for idx, (voluntario, qtd) in enumerate(voluntarios_ordenados):
        r_num = idx + 4
        ws_resumo.cell(row=r_num, column=1, value=voluntario).border = border_style
        cell_qtd = ws_resumo.cell(row=r_num, column=2, value=qtd)
        cell_qtd.border = border_style
        cell_qtd.alignment = center_align
        
    ws_resumo.column_dimensions["A"].width = 30
    ws_resumo.column_dimensions["B"].width = 18
    
    ws_resumo["D1"] = "TEXTO FORMATADO (WHATSAPP)"
    ws_resumo["D1"].font = Font(name="Calibri", size=14, bold=True, color="1B5E20")
    
    texto_completo_str = "\n".join(lista_texto_final)
    
    ws_resumo.merge_cells("D3:G25")
    cell_texto = ws_resumo["D3"]
    cell_texto.value = texto_completo_str
    cell_texto.alignment = Alignment(vertical="top", wrap_text=True)
    cell_texto.font = Font(name="Consolas", size=11, color="333333")
    cell_texto.fill = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")
    
    for r in range(3, 26):
        for c in range(4, 8):
            ws_resumo.cell(row=r, column=c).border = border_style

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


# --- CONFIGURAÇÃO DA INTERFACE STREAMLIT ---
st.set_page_config(page_title="Gerador Avançado de Escala", page_icon="🗓️", layout="wide")

st.title("🗓️ Gerador de Escala Dinâmico")
st.markdown("Gerencie voluntários, crie ou altere horários de reuniões livres na barra lateral e monte metas sob medida.")

# Inicialização do Session State
if "dados_escala" not in st.session_state:
    st.session_state.dados_escala = None
if "contagem_escalados" not in st.session_state:
    st.session_state.contagem_escalados = None
if "lista_texto_final" not in st.session_state:
    st.session_state.lista_texto_final = None

# --- GERENCIAMENTO DINÂMICO DE HORÁRIOS ---
if "lista_horarios" not in st.session_state:
    st.session_state.lista_horarios = ["9:30h", "11:30h", "17:30h"]

# Sidebar - Configurações Iniciais
st.sidebar.header("🛠️ 1. Configurações Iniciais")

nomes_input = st.sidebar.text_area(
    "Nomes dos Voluntários (com ou sem limites):",
    placeholder="Tiago Chaves(3), Maria Oliveira, Pedro Santos(1)",
    help="Adicione (X) ao lado do nome para forçar o voluntário a servir X vezes no mês."
)

# Processamento de voluntários
lista_voluntarios = []
limites_voluntarios = {}
tem_limite_definido = {}

itens_nomes = [n.strip() for n in nomes_input.split(",") if n.strip()]
for item in itens_nomes:
    match = re.search(r"^(.*?)\((\d+)\)$", item)
    if match:
        nome_limpo = match.group(1).strip()
        limite_num = int(match.group(2))
        lista_voluntarios.append(nome_limpo)
        limites_voluntarios[nome_limpo] = limite_num
        tem_limite_definido[nome_limpo] = True
    else:
        lista_voluntarios.append(item)
        limites_voluntarios[item] = 999
        tem_limite_definido[item] = False

# Painel de Edição de Horários na Barra Lateral
st.sidebar.markdown("---")
st.sidebar.header("⏰ 2. Configurar Horários")

# 1. Adicionar Horário
novo_horario = st.sidebar.text_input("Adicionar novo horário:", placeholder="ex: 19:00h")
if st.sidebar.button("➕ Adicionar"):
    if novo_horario and novo_horario.strip() not in st.session_state.lista_horarios:
        st.session_state.lista_horarios.append(novo_horario.strip())
        st.toast(f"Horário {novo_horario} adicionado!")
        st.rerun()

# 2. Remover ou Editar da Lista Existente
st.sidebar.markdown("**Horários Ativos atualmente:**")
horarios_remover = []

# Exibe os horários atuais como editores rápidos de texto ou remoção
for idx, h_atual in enumerate(st.session_state.lista_horarios):
    col_h_txt, col_h_del = st.sidebar.columns([4, 1])
    with col_h_txt:
        # Permite edição direta do texto do horário
        texto_editado = st.text_input(f"Editar #{idx+1}", value=h_atual, key=f"edit_h_{idx}", label_visibility="collapsed")
        if texto_editado != h_atual and texto_editado.strip():
            st.session_state.lista_horarios[idx] = texto_editado.strip()
            st.rerun()
    with col_h_del:
        if st.button("🗑️", key=f"del_h_{idx}"):
            horarios_remover.append(h_atual)

if horarios_remover:
    st.session_state.lista_horarios = [h for h in st.session_state.lista_horarios if h not in horarios_remover]
    st.rerun()


# --- CONFIGURAÇÃO DE PERÍODOS ---
st.sidebar.markdown("---")
st.sidebar.header("📅 3. Período da Escala")

meses_nomes = [
    "Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho",
    "Julho", "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro"
]
ano_atual = datetime.now().year

mes_selecionado_nome = st.sidebar.selectbox("Mês:", meses_nomes, index=datetime.now().month - 1)
mes_selecionado_num = meses_nomes.index(mes_selecionado_nome) + 1
ano_selecionado = st.sidebar.number_input("Ano:", min_value=ano_atual, max_value=ano_atual + 2, value=ano_atual)

domingos_do_mes = obter_domingos(mes_selecionado_num, ano_selecionado)
opcoes_domingos = [f"Domingo, dia {d['dia']} ({d['data_str']})" for d in domingos_do_mes]

domingos_escolhidos_texto = st.sidebar.multiselect("Quais domingos incluir?", options=opcoes_domingos, default=opcoes_domingos)

datas_finais = []
for d in domingos_do_mes:
    texto_corresponding = f"Domingo, dia {d['dia']} ({d['data_str']})"
    if texto_corresponding in domingos_escolhidos_texto:
        datas_finais.append(d['data_str'])


# --- CORPO PRINCIPAL: AJUSTE DE METAS INDIVIDUAIS ---
st.subheader("👥 Ajustar Meta de Voluntários por Reunião")
st.markdown("Defina abaixo o número necessário de pessoas. As colunas mudam se você alterar os horários na barra lateral.")

horarios_labels = st.session_state.lista_horarios

if not horarios_labels:
    st.warning("⚠️ Adicione pelo menos um horário na barra lateral para estruturar a escala!")
elif datas_finais:
    # Cria os dados base dinamicamente com base nos horários ativos
    dados_metas_padrao = []
    for data in datas_finais:
        linha_meta = {"DATA": data}
        for h in horarios_labels:
            linha_meta[h] = 1 # Valor padrão de 1 voluntário por horário
        dados_metas_padrao.append(linha_meta)
        
    df_metas_inicial = pd.DataFrame(dados_metas_padrao)
    df_metas_editado = st.data_editor(df_metas_inicial, disabled=["DATA"], hide_index=True, use_container_width=True)
else:
    st.info("💡 Escolha os domingos na barra lateral para abrir a tabela de metas por horário.")


# --- PROCESSAMENTO DO SORTEIO ---
st.markdown("---")
col_btn1, col_btn2 = st.columns([1, 5])

with col_btn1:
    gerar_clicado = st.button("🚀 Gerar Escala", type="primary")

with col_btn2:
    if st.button("🔄 Limpar / Resetar Resultados"):
        st.session_state.dados_escala = None
        st.session_state.contagem_escalados = None
        st.session_state.lista_texto_final = None
        st.rerun()

if gerar_clicado:
    if not lista_voluntarios:
        st.error("❌ Erro: Insira pelo menos o nome de um voluntário na barra lateral!")
    elif not datas_finais:
        st.error("❌ Erro: Selecione pelo menos um domingo na barra lateral!")
    elif not horarios_labels:
        st.error("❌ Erro: Configure ao menos um horário de reunião na barra lateral!")
    else:
        contagem_escalados = {nome: 0 for nome in lista_voluntarios}
        dados_escala = []
        lista_texto_final = [
            "✨ ESCALA FINALIZADA ✨",
            f"Período: {mes_selecionado_nome} de {ano_selecionado}",
            "Confira abaixo sua escala e horários:",
            ""
        ]

        dict_metas = df_metas_editado.set_index("DATA").to_dict(orient="index")

        # Algoritmo de dupla esteira com colunas totalmente dinâmicas
        for data_evento in datas_finais:
            row_escala = {"DATA": data_evento}
            lista_texto_final.append(f"📅 DATA: {data_evento}")

            for horario in horarios_labels:
                escolhidos_do_horario = []
                qtd_vagas_do_horario = int(dict_metas[data_evento].get(horario, 0))

                for _ in range(qtd_vagas_do_horario):
                    validos = [
                        n for n in lista_voluntarios 
                        if n not in escolhidos_do_horario and contagem_escalados[n] < limites_voluntarios[n]
                    ]

                    if validos:
                        prioritarios_com_meta = [
                            n for n in validos 
                            if tem_limite_definido[n] and contagem_escalados[n] < limites_voluntarios[n]
                        ]

                        if prioritarios_com_meta:
                            min_p = min(contagem_escalados[n] for n in prioritarios_com_meta)
                            candidatos = [n for n in prioritarios_com_meta if contagem_escalados[n] == min_p]
                        else:
                            min_p = min(contagem_escalados[n] for n in validos)
                            candidatos = [n for n in validos if contagem_escalados[n] == min_p]

                        escolhido = random.choice(candidatos)
                        contagem_escalados[escolhido] += 1
                        escolhidos_do_horario.append(escolhido)

                if escolhidos_do_horario:
                    string_escolhidos = ", ".join(escolhidos_do_horario)
                    if len(escolhidos_do_horario) < qtd_vagas_do_horario:
                        string_escolhidos += " (⚠️ VAGA INCOMPLETA)"
                    row_escala[horario] = string_escolhidos
                    lista_texto_final.append(f"   ⏰ {horario} → {string_escolhidos}")
                else:
                    row_escala[horario] = "VAGO"
                    lista_texto_final.append(f"   ⏰ {horario} → ⚠️ VAGO")

            dados_escala.append(row_escala)

        st.session_state.dados_escala = dados_escala
        st.session_state.contagem_escalados = contagem_escalados
        st.session_state.lista_texto_final = lista_texto_final


# --- RENDERIZAÇÃO DO RESULTADO ---
if st.session_state.dados_escala is not None:
    st.success(f"🎉 Escala de {mes_selecionado_nome} calculada com sucesso!")
    col1, col2 = st.columns([2, 1])

    with col1:
        st.subheader("📊 Visualização da Escala Resultante")
        df_escala = pd.DataFrame(st.session_state.dados_escala)
        st.dataframe(df_escala, use_container_width=True, hide_index=True)

        excel_bytes = exportar_excel_completo(
            st.session_state.dados_escala, st.session_state.contagem_escalados, st.session_state.lista_texto_final, horarios_labels
        )
        st.download_button(
            label="📥 Baixar Documento Completo (Excel)",
            data=excel_bytes,
            file_name=f"Escala_Completa_{mes_selecionado_nome}_{ano_selecionado}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        with st.expander("📊 Métrica Rápida na Tela (Quem trabalhou quanto?)"):
            df_metricas = pd.DataFrame(list(st.session_state.contagem_escalados.items()), columns=["Voluntário", "Vezes Escalado"])
            st.dataframe(df_metricas.sort_values(by="Vezes Escalado", ascending=False), hide_index=True)

    with col2:
        st.subheader("💬 Texto Pronto para Redes Sociais")
        texto_completo = "\n".join(st.session_state.lista_texto_final)
        st.text_area("Copie o conteúdo abaixo:", value=texto_completo, height=380)

        st.download_button(
            label="📄 Baixar Lista em TXT", data=texto_completo, file_name=f"Texto_Escala_{mes_selecionado_nome}.txt", mime="text/plain"
        )